import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def set_shape_text(shape, new_text):
    """Sets text on a single-line or simple shape while preserving styling."""
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    if not tf.paragraphs:
        p = tf.add_paragraph()
    else:
        p = tf.paragraphs[0]
        # remove extra paragraphs if any
        while len(tf.paragraphs) > 1:
            p_elem = tf.paragraphs[-1]._p
            p_elem.getparent().remove(p_elem)
            
    if not p.runs:
        r = p.add_run()
        r.text = new_text
    else:
        # Preserve first run font attributes
        r0 = p.runs[0]
        font_name = r0.font.name
        font_size = r0.font.size
        font_bold = r0.font.bold
        font_color = None
        if r0.font.color and r0.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB:
            font_color = r0.font.color.rgb
            
        r0.text = new_text
        if font_name: r0.font.name = font_name
        if font_size: r0.font.size = font_size
        if font_bold is not None: r0.font.bold = font_bold
        if font_color: r0.font.color.rgb = font_color
        
        # Remove any extra runs
        while len(p.runs) > 1:
            r_elem = p.runs[-1]._r
            r_elem.getparent().remove(r_elem)

def set_shape_paragraphs(shape, paragraphs_list, default_font_size_pt=10):
    """Sets multiple paragraphs/bullets in a shape preserving style."""
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    # Sample style from first run
    font_name = "Inter"
    font_color = RGBColor(18, 87, 83)
    if tf.paragraphs and tf.paragraphs[0].runs:
        r0 = tf.paragraphs[0].runs[0]
        if r0.font.name: font_name = r0.font.name
        if r0.font.color and r0.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB:
            font_color = r0.font.color.rgb

    # Clear paragraphs
    while len(tf.paragraphs) > 1:
        p_elem = tf.paragraphs[-1]._p
        p_elem.getparent().remove(p_elem)
        
    p0 = tf.paragraphs[0]
    p0.text = ""
    
    for i, item in enumerate(paragraphs_list):
        if i == 0:
            p = p0
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = item
        r.font.name = font_name
        r.font.size = Pt(default_font_size_pt)
        r.font.color.rgb = font_color

def get_shape_by_name(slide, name):
    for s in slide.shapes:
        if s.name == name:
            return s
    return None

def enhance_deck(src_path, dst_path):
    prs = pptx.Presentation(src_path)
    print("Loaded presentation with", len(prs.slides), "slides")

    # ==================== SLIDE 1 ====================
    s1 = prs.slides[0]
    set_shape_text(get_shape_by_name(s1, "Text 1"), "GOOGLE BUILD WITH AI HACKATHON 2.0")
    set_shape_text(get_shape_by_name(s1, "Text 2"), "KisanSathi")
    set_shape_text(get_shape_by_name(s1, "Text 3"), "AI-Powered Agricultural Market Intelligence & Farmer Support Platform")
    set_shape_text(get_shape_by_name(s1, "Text 4"), "Full-Stack AI Agricultural Platform · Google Cloud Hack2skill")
    set_shape_text(get_shape_by_name(s1, "Text 7"), "APMC Mandi Intelligence")
    set_shape_text(get_shape_by_name(s1, "Text 10"), "Gemini 2.0 Agronomy")
    set_shape_text(get_shape_by_name(s1, "Text 13"), "Sentinel-2 Geospatial")
    set_shape_text(get_shape_by_name(s1, "Text 16"), "10 Regional Languages")
    print("Slide 1 updated.")

    # ==================== SLIDE 2 ====================
    s2 = prs.slides[1]
    set_shape_text(get_shape_by_name(s2, "Text 1"), "THE PROBLEM")
    set_shape_text(get_shape_by_name(s2, "Text 2"), "Small Farmers Lack Unified Decision Support")
    set_shape_text(get_shape_by_name(s2, "Text 3"), "Fragmented Agricultural Information")
    set_shape_text(get_shape_by_name(s2, "Text 4"), "Market prices, weather forecasts, and soil advisory scattered across disconnected government portals.")
    set_shape_text(get_shape_by_name(s2, "Text 5"), "Severe Price Asymmetry & Middlemen Opacity")
    set_shape_text(get_shape_by_name(s2, "Text 6"), "Price opacity forces distress sales to intermediaries without modal, min, and max mandi price transparency.")
    set_shape_text(get_shape_by_name(s2, "Text 7"), "Zero Pre-Harvest Value Estimation")
    set_shape_text(get_shape_by_name(s2, "Text 8"), "Farmers struggle to calculate net returns after factoring in packaging, transport logistics, and mandi fees.")
    set_shape_text(get_shape_by_name(s2, "Text 9"), "Delayed Field Crop Disease Diagnosis")
    set_shape_text(get_shape_by_name(s2, "Text 10"), "Preventable crop failures spread rapidly with no accessible, on-field AI visual diagnosis tool for leaves.")
    set_shape_text(get_shape_by_name(s2, "Text 11"), "Language & Digital Literacy Barriers")
    set_shape_text(get_shape_by_name(s2, "Text 12"), "Scientific advisories and market analytics are locked behind complex English-only dashboards.")
    print("Slide 2 updated.")

    # ==================== SLIDE 3 ====================
    s3 = prs.slides[2]
    set_shape_text(get_shape_by_name(s3, "Text 1"), "OUR SOLUTION")
    set_shape_text(get_shape_by_name(s3, "Text 2"), "KisanSathi — One Platform, All Decisions")
    set_shape_text(get_shape_by_name(s3, "Text 3"), "A unified farmer intelligence ecosystem combining real-time APMC mandi analytics, Gemini 2.0 multimodal agronomy, Sentinel-2 geospatial indices, and official cadastre.")
    set_shape_text(get_shape_by_name(s3, "Text 4"), "Mandi Market Intelligence")
    set_shape_text(get_shape_by_name(s3, "Text 5"), "Real-time modal rates across 3,000+ mandis")
    set_shape_text(get_shape_by_name(s3, "Text 6"), "AI Explanation & Advisory")
    set_shape_text(get_shape_by_name(s3, "Text 7"), "Gemini 2.0 market insights & seasonal crop guide")
    set_shape_text(get_shape_by_name(s3, "Text 8"), "Satellite & NDVI Telemetry")
    set_shape_text(get_shape_by_name(s3, "Text 9"), "Sentinel-2 & GEE crop vigor & water stress")
    set_shape_text(get_shape_by_name(s3, "Text 10"), "Hyperlocal Ag Weather")
    set_shape_text(get_shape_by_name(s3, "Text 11"), "7-day microclimate data & spray window alerts")
    set_shape_text(get_shape_by_name(s3, "Text 12"), "Plant Disease Diagnosis")
    set_shape_text(get_shape_by_name(s3, "Text 13"), "Gemini Vision leaf pathology & organic remedies")
    set_shape_text(get_shape_by_name(s3, "Text 14"), "10 Languages & IVR Support")
    set_shape_text(get_shape_by_name(s3, "Text 15"), "Multilingual UI + 1800-180-1551 Kisan Call Centre")
    print("Slide 3 updated.")

    # ==================== SLIDE 4 ====================
    s4 = prs.slides[3]
    set_shape_text(get_shape_by_name(s4, "Text 1"), "KEY FEATURES")
    set_shape_text(get_shape_by_name(s4, "Text 2"), "10 Production Feature Modules")
    set_shape_text(get_shape_by_name(s4, "Text 5"), "Mandi / Market Intelligence")
    set_shape_text(get_shape_by_name(s4, "Text 6"), "Real-time APMC price lookup across 3,000+ mandis with modal spread analysis")
    set_shape_text(get_shape_by_name(s4, "Text 9"), "Price Trend Analysis")
    set_shape_text(get_shape_by_name(s4, "Text 10"), "30-day price trends, volatility metrics & arrival volume tracking")
    set_shape_text(get_shape_by_name(s4, "Text 13"), "AI Market Explanation")
    set_shape_text(get_shape_by_name(s4, "Text 14"), "Gemini 2.0 grounded market context in simple regional dialects")
    set_shape_text(get_shape_by_name(s4, "Text 17"), "Contextual Crop Advisory")
    set_shape_text(get_shape_by_name(s4, "Text 18"), "Agro-climatic crop guidance, sowing windows & nutrient schedules")
    set_shape_text(get_shape_by_name(s4, "Text 21"), "Plant Disease Diagnosis")
    set_shape_text(get_shape_by_name(s4, "Text 22"), "Upload leaf photo → Gemini Vision diagnoses disease & organic remedies")
    set_shape_text(get_shape_by_name(s4, "Text 25"), "Satellite & NDVI Telemetry")
    set_shape_text(get_shape_by_name(s4, "Text 26"), "Google Earth Engine + Sentinel-2 multispectral field vigor mapping")
    set_shape_text(get_shape_by_name(s4, "Text 29"), "Hyperlocal Ag Weather")
    set_shape_text(get_shape_by_name(s4, "Text 30"), "Open-Meteo 7-day agricultural forecasts with optimal spray windows")
    set_shape_text(get_shape_by_name(s4, "Text 33"), "Government Cadastre & Schemes")
    set_shape_text(get_shape_by_name(s4, "Text 34"), "AgriStack GID, ISRO Bhuvan, PM-KISAN, PMFBY & State RoR integrations")
    set_shape_text(get_shape_by_name(s4, "Text 37"), "Field Polygon Sculpting")
    set_shape_text(get_shape_by_name(s4, "Text 38"), "Trace parcel boundaries with real-time geodesic area calculation (ha/ac)")
    set_shape_text(get_shape_by_name(s4, "Text 41"), "10 Languages & Offline Sync")
    set_shape_text(get_shape_by_name(s4, "Text 42"), "10 Indian languages, low-bandwidth caching & recent-search resilience")
    print("Slide 4 updated.")

    # ==================== SLIDE 5 ====================
    s5 = prs.slides[4]
    set_shape_text(get_shape_by_name(s5, "Text 1"), "HOW IT WORKS")
    set_shape_text(get_shape_by_name(s5, "Text 2"), "End-to-End Processing Workflow")
    set_shape_text(get_shape_by_name(s5, "Text 9"), "✅ Verified Structured Data")
    set_shape_text(get_shape_by_name(s5, "Text 10"), "Official mandi prices from data.gov.in / Agmarknet, Open-Meteo & Sentinel-2")
    set_shape_text(get_shape_by_name(s5, "Text 12"), "🔢 Deterministic Calculations")
    set_shape_text(get_shape_by_name(s5, "Text 13"), "Gross value, net return, transport deductions & trend stats — zero AI hallucination")
    set_shape_text(get_shape_by_name(s5, "Text 15"), "🤖 Grounded Gemini 2.0 AI")
    set_shape_text(get_shape_by_name(s5, "Text 16"), "Translates, explains & diagnoses — strictly bounded by verified calculations")
    print("Slide 5 updated.")

    # ==================== SLIDE 6 (ARCHITECTURE) ====================
    s6 = prs.slides[5]
    # Move headers to top
    s0 = get_shape_by_name(s6, "Shape 0")
    img0 = get_shape_by_name(s6, "Image 0")
    t1 = get_shape_by_name(s6, "Text 1")
    t2 = get_shape_by_name(s6, "Text 2")
    
    if s0: s0.top = Inches(0.65)
    if img0: img0.top = Inches(0.71)
    if t1:
        t1.top = Inches(0.70)
        set_shape_text(t1, "SYSTEM ARCHITECTURE")
    if t2:
        t2.top = Inches(0.98)
        t2.height = Inches(0.40)
        set_shape_text(t2, "Multi-Tier Cloud, AI & Data Architecture")

    # Add subtitle text box
    sub_box = s6.shapes.add_textbox(Inches(0.72), Inches(1.42), Inches(11.89), Inches(0.35))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Complete structural separation between verified deterministic engines, Gemini 2.0 AI, and agricultural data mesh."
    p_sub.font.name = "Inter"
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = RGBColor(18, 87, 83)

    # Remove broken placeholder Image 1
    for shape in list(s6.shapes):
        if shape.name == "Image 1":
            sp = shape._element
            sp.getparent().remove(sp)

    # Build 4 clean Architecture Tiers
    # Slide width: 13.33 in, margins: 0.72 in each side -> available width: 11.89 in
    tiers_data = [
        {
            "tier_name": "TIER 1: CLIENT PRESENTATION LAYER (React 19 SPA)",
            "y": 1.85, "h": 1.05,
            "bg_color": RGBColor(244, 248, 246),
            "border_color": RGBColor(212, 247, 246),
            "items": [
                ("Farmer Responsive Web App", "React 19, TypeScript, Vite & Tailwind CSS with high-contrast UI"),
                ("Interactive Geospatial GIS", "Leaflet maps with real-time geodesic polygon sculpting (ha/ac)"),
                ("10 Regional Indian Languages", "Hindi, Kannada, Telugu, Tamil, Marathi, Bengali, Gujarati, etc."),
                ("Voice & Offline Resilience", "Speech synthesis, Kisan Call Centre IVR scripts & cached queries")
            ]
        },
        {
            "tier_name": "TIER 2: API GATEWAY & APPLICATION SERVER (Node.js 20+ Express)",
            "y": 3.02, "h": 1.05,
            "bg_color": RGBColor(244, 248, 246),
            "border_color": RGBColor(212, 247, 246),
            "items": [
                ("RESTful Microservices Gateway", "30+ secured modular endpoints routing farmer analytics requests"),
                ("Security & Rate Limiting", "Input sanitization, origin verification & DDOS request throttling"),
                ("In-Memory Hot Caching", "Sub-second TTL caching for mandi commodity price records"),
                ("Cadastre Query Orchestration", "AgriStack GID, ISRO Bhuvan & State land record RoR dispatch")
            ]
        },
        {
            "tier_name": "TIER 3: CORE PROCESSING & DUAL INTELLIGENCE ENGINES",
            "y": 4.19, "h": 1.15,
            "bg_color": RGBColor(236, 246, 240),
            "border_color": RGBColor(18, 87, 83),
            "items": [
                ("Deterministic Engine (Zero AI Math)", "Exact modal price spreads, gross revenue & net farmer returns"),
                ("Google Gemini 2.0 AI Core", "Multimodal Vision leaf diagnosis & grounded regional narratives"),
                ("Geospatial & Telemetry Core", "Google Earth Engine + Sentinel-2 multispectral NDVI/NDWI indices"),
                ("Zero-Hallucination Guardrail", "Strict boundary isolating financial arithmetic from LLM text")
            ]
        },
        {
            "tier_name": "TIER 4: FEDERATED DATA MESH & PERSISTENCE LAYER",
            "y": 5.46, "h": 1.05,
            "bg_color": RGBColor(244, 248, 246),
            "border_color": RGBColor(212, 247, 246),
            "items": [
                ("data.gov.in / Agmarknet", "Daily verified modal, min & max prices across 3,000+ mandis"),
                ("Open-Meteo & IMD Ag-Weather", "Hyperlocal 7-day temperature, rainfall, wind & spray windows"),
                ("AgriStack & ISRO Bhuvan", "Official national land parcel boundaries & soil health cards"),
                ("Cloud Firestore & ETSI NGSI-LD", "Secure persistence store & federated Digital Public Good exchange")
            ]
        }
    ]

    for t_idx, tier in enumerate(tiers_data):
        ty = Inches(tier["y"])
        th = Inches(tier["h"])
        tw = Inches(11.89)
        tx = Inches(0.72)

        # Outer Tier Container
        outer = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, tx, ty, tw, th)
        outer.fill.solid()
        outer.fill.fore_color.rgb = tier["bg_color"]
        outer.line.color.rgb = tier["border_color"]
        outer.line.width = Pt(1.2)

        # Tier Title Pill
        title_box = s6.shapes.add_textbox(tx + Inches(0.12), ty + Inches(0.04), tw - Inches(0.24), Inches(0.25))
        tf_t = title_box.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = tier["tier_name"]
        p_t.font.name = "Manrope SemiBold"
        p_t.font.size = Pt(9.5)
        p_t.font.bold = True
        p_t.font.color.rgb = RGBColor(18, 87, 83)

        # 4 Sub-blocks inside tier
        sub_count = len(tier["items"])
        gap = Inches(0.12)
        total_sub_w = tw - Inches(0.24) - gap * (sub_count - 1)
        sub_w = total_sub_w / sub_count
        sub_h = th - Inches(0.36)
        sub_y = ty + Inches(0.28)

        for s_i, (head, desc) in enumerate(tier["items"]):
            sub_x = tx + Inches(0.12) + s_i * (sub_w + gap)
            sub_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sub_x, sub_y, sub_w, sub_h)
            sub_card.fill.solid()
            sub_card.fill.fore_color.rgb = RGBColor(255, 255, 255)
            sub_card.line.color.rgb = RGBColor(220, 230, 226)
            sub_card.line.width = Pt(0.75)

            # Sub-card text
            tf_c = sub_card.text_frame
            tf_c.word_wrap = True
            tf_c.vertical_anchor = MSO_ANCHOR.TOP
            tf_c.margin_left = Inches(0.08)
            tf_c.margin_right = Inches(0.08)
            tf_c.margin_top = Inches(0.06)
            tf_c.margin_bottom = Inches(0.04)

            p_h = tf_c.paragraphs[0]
            p_h.text = head
            p_h.font.name = "Manrope SemiBold"
            p_h.font.size = Pt(8.5)
            p_h.font.bold = True
            p_h.font.color.rgb = RGBColor(18, 87, 83)
            p_h.space_after = Pt(2)

            p_d = tf_c.add_paragraph()
            p_d.text = desc
            p_d.font.name = "Inter"
            p_d.font.size = Pt(7.5)
            p_d.font.color.rgb = RGBColor(70, 95, 88)

    # Add connector badges between tiers
    conn_y_positions = [2.90, 4.07, 5.34]
    for cy in conn_y_positions:
        conn = s6.shapes.add_textbox(Inches(6.3), Inches(cy), Inches(0.7), Inches(0.2))
        tf_conn = conn.text_frame
        p_c = tf_conn.paragraphs[0]
        p_c.text = "▼"
        p_c.alignment = PP_ALIGN.CENTER
        p_c.font.name = "Inter"
        p_c.font.size = Pt(7)
        p_c.font.color.rgb = RGBColor(18, 87, 83)

    print("Slide 6 updated with comprehensive 4-tier architecture.")

    # ==================== SLIDE 7 ====================
    s7 = prs.slides[6]
    set_shape_text(get_shape_by_name(s7, "Text 1"), "TECHNOLOGY STACK")
    set_shape_text(get_shape_by_name(s7, "Text 2"), "Built With Modern Technologies")
    set_shape_text(get_shape_by_name(s7, "Text 3"), "Frontend Application")
    set_shape_text(get_shape_by_name(s7, "Text 4"), "React 19 · TypeScript · Vite · Tailwind CSS · Framer Motion · Leaflet GIS")
    set_shape_text(get_shape_by_name(s7, "Text 5"), "Backend & API Gateway")
    set_shape_text(get_shape_by_name(s7, "Text 6"), "Node.js 20+ · Express.js · RESTful Gateway · 30+ endpoints · In-Memory Cache")
    set_shape_text(get_shape_by_name(s7, "Text 7"), "Artificial Intelligence")
    set_shape_text(get_shape_by_name(s7, "Text 8"), "Google Gemini 2.0 Flash · Gemini Multimodal Vision · Grounded Prompts")
    set_shape_text(get_shape_by_name(s7, "Text 9"), "Deployment & Cloud")
    set_shape_text(get_shape_by_name(s7, "Text 10"), "Vercel Serverless · Google Cloud Run · Docker Containerization")
    set_shape_text(get_shape_by_name(s7, "Text 11"), "Satellite & Geospatial")
    set_shape_text(get_shape_by_name(s7, "Text 12"), "Google Earth Engine · Sentinel-2 Multispectral · Geodesic Polygon Tools")
    set_shape_text(get_shape_by_name(s7, "Text 13"), "Database & Storage")
    set_shape_text(get_shape_by_name(s7, "Text 14"), "Google Cloud Firestore · In-Memory Hot Store · Local Storage Resilience")
    set_shape_text(get_shape_by_name(s7, "Text 15"), "Government & Open Standards")
    set_shape_text(get_shape_by_name(s7, "Text 16"), "data.gov.in (Agmarknet) · Open-Meteo · AgriStack · ISRO Bhuvan · ETSI NGSI-LD")
    print("Slide 7 updated.")

    # ==================== SLIDE 8 ====================
    s8 = prs.slides[7]
    set_shape_text(get_shape_by_name(s8, "Text 1"), "RESPONSIBLE AI ARCHITECTURE")
    set_shape_text(get_shape_by_name(s8, "Text 2"), "The Zero-Hallucination Triad: AI + Deterministic Rigor")
    set_shape_text(get_shape_by_name(s8, "Text 3"), "AI explains and assists — it never invents market prices or financial calculations.")
    set_shape_text(get_shape_by_name(s8, "Text 5"), "✅ Verified Data Layer")
    set_shape_paragraphs(get_shape_by_name(s8, "Text 6"), [
        "Real mandi / APMC prices across 3,000+ mandis",
        "Modal, minimum & maximum price benchmarks",
        "Arrival volumes & seasonal historical records",
        "Official open data: data.gov.in / Agmarknet",
        "Hyperlocal weather via Open-Meteo & IMD"
    ], default_font_size_pt=9.5)

    set_shape_text(get_shape_by_name(s8, "Text 7"), "🔢 Deterministic Processing")
    set_shape_paragraphs(get_shape_by_name(s8, "Text 8"), [
        "Yield & quantity normalization (Quintals/Kg)",
        "Gross harvest value & net-return calculations",
        "Transport, packaging & commission deductions",
        "Trend statistics & volatility index computation",
        "Zero-LLM math guarantee prevents price fabrication"
    ], default_font_size_pt=9.5)

    set_shape_text(get_shape_by_name(s8, "Text 10"), "🤖 Grounded Gemini 2.0 AI")
    set_shape_paragraphs(get_shape_by_name(s8, "Text 11"), [
        "Natural-language market price explanation",
        "10 Indian languages (Hindi, Kannada, Telugu, etc.)",
        "Agro-climatic crop & irrigation guidance",
        "Leaf disease diagnosis via Gemini Vision",
        "Kisan Call Centre 1800-180-1551 IVR scripts"
    ], default_font_size_pt=9.5)
    print("Slide 8 updated.")

    # ==================== SLIDE 9 ====================
    s9 = prs.slides[8]
    set_shape_text(get_shape_by_name(s9, "Text 1"), "USER JOURNEY")
    set_shape_text(get_shape_by_name(s9, "Text 2"), "End-to-End 6-Step Farmer Flow")
    set_shape_text(get_shape_by_name(s9, "Text 4"), "Select Language & Crop")
    set_shape_text(get_shape_by_name(s9, "Text 5"), "Choose from 10 Indian languages (English, Hindi, Kannada, Telugu, etc.)")
    set_shape_text(get_shape_by_name(s9, "Text 7"), "Search APMC Mandi Rates")
    set_shape_text(get_shape_by_name(s9, "Text 8"), "KisanSathi retrieves structured live modal prices from official Agmarknet data")
    set_shape_text(get_shape_by_name(s9, "Text 10"), "Net-Return & Value Calculation")
    set_shape_text(get_shape_by_name(s9, "Text 11"), "Deterministic calculation of gross yield value and net returns after transport")
    set_shape_text(get_shape_by_name(s9, "Text 13"), "Grounded AI Market Insights")
    set_shape_text(get_shape_by_name(s9, "Text 14"), "Gemini translates complex market spreads into clear regional-language advice")
    set_shape_text(get_shape_by_name(s9, "Text 16"), "Field Health & Disease Diagnosis")
    set_shape_text(get_shape_by_name(s9, "Text 17"), "Trace parcel polygon on Sentinel-2 satellite map & photo leaf for Gemini diagnosis")
    set_shape_text(get_shape_by_name(s9, "Text 19"), "Checklist, Decision & IVR Support")
    set_shape_text(get_shape_by_name(s9, "Text 20"), "Review pre-sale checklist or trigger 1-tap Kisan Call Centre (1800-180-1551) escalation")
    print("Slide 9 updated.")

    # ==================== SLIDE 10 ====================
    s10 = prs.slides[9]
    set_shape_text(get_shape_by_name(s10, "Text 1"), "TECHNICAL HIGHLIGHTS")
    set_shape_text(get_shape_by_name(s10, "Text 2"), "What Makes KisanSathi Different")
    set_shape_text(get_shape_by_name(s10, "Text 3"), "Dual-Engine Hybrid Rigor")
    set_shape_text(get_shape_by_name(s10, "Text 4"), "Verified government data + deterministic math + AI synthesis — 0% hallucinated numbers")
    set_shape_text(get_shape_by_name(s10, "Text 5"), "Full-Spectrum Satellite GIS")
    set_shape_text(get_shape_by_name(s10, "Text 6"), "Google Earth Engine + Sentinel-2 multispectral NDVI/NDWI crop health mapping")
    set_shape_text(get_shape_by_name(s10, "Text 7"), "Interactive Polygon Sculpting")
    set_shape_text(get_shape_by_name(s10, "Text 8"), "Real-time geodesic field area calculation (ha/ac) with multi-cadastre overlays")
    set_shape_text(get_shape_by_name(s10, "Text 9"), "Production Cloud-Ready")
    set_shape_text(get_shape_by_name(s10, "Text 10"), "Vercel + Google Cloud Run + Docker with 30+ secured REST endpoints")

    set_shape_text(get_shape_by_name(s10, "Text 12"), "FUTURE SCOPE")
    set_shape_text(get_shape_by_name(s10, "Text 13"), "Roadmap & Expansion")
    set_shape_text(get_shape_by_name(s10, "Text 14"), "AgriStack Farmer ID Locker")
    set_shape_text(get_shape_by_name(s10, "Text 15"), "Secure farmer digital identity & personalized soil health record vault")
    set_shape_text(get_shape_by_name(s10, "Text 16"), "WhatsApp & IVR Voice Bot")
    set_shape_text(get_shape_by_name(s10, "Text 17"), "Voice-based market inquiries on feature phones via IVR & WhatsApp bot")
    set_shape_text(get_shape_by_name(s10, "Text 18"), "Cooperative Mandi Logistics")
    set_shape_text(get_shape_by_name(s10, "Text 19"), "Collective transport pooling for smallholders to slash freight costs")
    set_shape_text(get_shape_by_name(s10, "Text 20"), "Soil Carbon Credit Verification")
    set_shape_text(get_shape_by_name(s10, "Text 21"), "Satellite-verified regenerative practices for farmer carbon credit monetization")

    set_shape_text(get_shape_by_name(s10, "Text 23"), "KisanSathi — From Agricultural Data to Better Farmer Decisions")
    set_shape_text(get_shape_by_name(s10, "Text 25"), "📊 Data-Driven")
    set_shape_text(get_shape_by_name(s10, "Text 26"), "Verified market, satellite and cadastre data across 3,000+ mandis")
    set_shape_text(get_shape_by_name(s10, "Text 28"), "🤖 Responsible AI")
    set_shape_text(get_shape_by_name(s10, "Text 29"), "Gemini 2.0 explains, advises and diagnoses with zero hallucination")
    set_shape_text(get_shape_by_name(s10, "Text 31"), "🌾 Farmer-First")
    set_shape_text(get_shape_by_name(s10, "Text 32"), "10 Indian languages, accessible UI & 1800-180-1551 Kisan Call Centre")
    set_shape_text(get_shape_by_name(s10, "Text 33"), "Thank You  ·  Questions & Demonstration")
    print("Slide 10 updated.")

    prs.save(dst_path)
    print("Successfully saved enhanced presentation to:", dst_path)

if __name__ == "__main__":
    src = "KisanSathi_Presentation.pptx"
    dst = "KisanSathi_Presentation.pptx"
    enhance_deck(src, dst)
